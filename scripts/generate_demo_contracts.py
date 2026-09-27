from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_ROOT = ROOT / "data" / "contracts"


@dataclass(frozen=True)
class ContractCase:
    filename: str
    category: str
    number: str
    city: str
    date: str
    employer: str
    employer_representative: str
    employee: str
    position: str
    department: str
    workplace: str | None
    start_date: str
    salary_clause: str
    payment_clause: str
    schedule_clause: str
    hours_clause: str
    leave_clause: str
    probation_clause: str
    function_clause: str
    structure_variant: int = 1


CASES = [
    ContractCase(
        filename="correct_01_accountant.docx",
        category="correct",
        number="01-2026",
        city="Москва",
        date="01 сентября 2026 года",
        employer="ООО Альфа Документ",
        employer_representative="генерального директора Орловой Марии Викторовны",
        employee="Иванова Ивана Ивановича",
        position="бухгалтера",
        department="финансовый отдел",
        workplace="г. Москва, ул. Деловая, д. 10",
        start_date="02 сентября 2026 года",
        salary_clause="Работнику устанавливается должностной оклад 85 000 рублей в месяц. Доплаты и премии выплачиваются в соответствии с положением об оплате труда.",
        payment_clause="Заработная плата выплачивается два раза в месяц: 20 числа текущего месяца и 5 числа следующего месяца путем перевода в кредитную организацию, указанную Работником.",
        schedule_clause="Работнику устанавливается пятидневная рабочая неделя с двумя выходными днями - субботой и воскресеньем. Рабочий день начинается в 9 часов 00 минут и заканчивается в 18 часов 00 минут. Перерыв для отдыха и питания предоставляется с 13 часов 00 минут до 14 часов 00 минут.",
        hours_clause="Нормальная продолжительность рабочего времени составляет 40 часов в неделю.",
        leave_clause="Работнику предоставляется ежегодный основной оплачиваемый отпуск продолжительностью 28 календарных дней.",
        probation_clause="Работнику устанавливается испытание сроком три месяца.",
        function_clause="Работник ведет бухгалтерский учет, оформляет первичные документы и готовит установленную отчетность.",
    ),
    ContractCase(
        filename="correct_02_business_analyst.docx",
        category="correct",
        number="02-2026",
        city="Казань",
        date="03 сентября 2026 года",
        employer="ООО Волга Системы",
        employer_representative="директора Соколова Павла Андреевича",
        employee="Петровой Анны Олеговны",
        position="бизнес-аналитика",
        department="отдел развития продуктов",
        workplace="г. Казань, ул. Университетская, д. 18",
        start_date="07 сентября 2026 года",
        salary_clause="Работнику устанавливается должностной оклад 120 000 рублей в месяц и ежеквартальная премия на условиях положения о премировании.",
        payment_clause="Заработная плата перечисляется на счет, выбранный Работником, каждые полмесяца: 22 числа текущего месяца и 7 числа следующего месяца.",
        schedule_clause="Режим работы: пятидневная рабочая неделя, выходные дни - суббота и воскресенье; начало работы в 8 часов 30 минут, окончание в 17 часов 30 минут; перерыв с 12 часов 30 минут до 13 часов 30 минут.",
        hours_clause="Продолжительность рабочего времени составляет 40 часов в неделю.",
        leave_clause="Ежегодный основной оплачиваемый отпуск составляет 28 календарных дней.",
        probation_clause="Работник принимается без испытания.",
        function_clause="Работник собирает и описывает требования к информационным системам, моделирует бизнес-процессы и согласовывает результаты анализа с заказчиками.",
        structure_variant=2,
    ),
    ContractCase(
        filename="correct_03_office_manager.docx",
        category="correct",
        number="03-2026",
        city="Екатеринбург",
        date="08 сентября 2026 года",
        employer="ООО Урал Консалт",
        employer_representative="генерального директора Кузнецова Ильи Романовича",
        employee="Смирновой Елены Дмитриевны",
        position="офис-менеджера",
        department="административный отдел",
        workplace="г. Екатеринбург, проспект Ленина, д. 41",
        start_date="09 сентября 2026 года",
        salary_clause="Должностной оклад Работника составляет 70 000 рублей в месяц. Размер и основания стимулирующих выплат определяются положением о премировании.",
        payment_clause="Работодатель выплачивает заработную плату 18 числа текущего месяца и 3 числа следующего месяца переводом на счет в банке, указанном Работником.",
        schedule_clause="Работник работает пять дней в неделю с понедельника по пятницу с 9 часов 30 минут до 18 часов 30 минут. Перерыв для отдыха и питания - один час с 13 часов 30 минут.",
        hours_clause="Продолжительность рабочей недели не превышает 40 часов.",
        leave_clause="Продолжительность ежегодного основного оплачиваемого отпуска - 28 календарных дней.",
        probation_clause="Стороны согласовали испытание продолжительностью два месяца.",
        function_clause="Работник организует документооборот офиса, прием посетителей, снабжение канцелярскими товарами и взаимодействие с поставщиками.",
        structure_variant=3,
    ),
    ContractCase(
        filename="correct_04_hr_specialist.docx",
        category="correct",
        number="04-2026",
        city="Новосибирск",
        date="10 сентября 2026 года",
        employer="ООО Сибирь Кадры",
        employer_representative="директора Морозовой Алины Сергеевны",
        employee="Волкова Никиты Максимовича",
        position="специалиста по кадрам",
        department="служба управления персоналом",
        workplace="г. Новосибирск, Красный проспект, д. 75",
        start_date="14 сентября 2026 года",
        salary_clause="За выполнение трудовой функции Работнику устанавливается оклад 78 000 рублей в месяц. Компенсационные и стимулирующие выплаты производятся по локальным нормативным актам Работодателя.",
        payment_clause="Заработная плата выплачивается 25 числа текущего месяца и 10 числа следующего месяца на банковский счет, который Работник указал в заявлении.",
        schedule_clause="Устанавливается пятидневная рабочая неделя с выходными в субботу и воскресенье. Время работы - с 8 часов 00 минут до 17 часов 00 минут, перерыв - с 12 часов 00 минут до 13 часов 00 минут.",
        hours_clause="Рабочее время составляет 40 часов в неделю.",
        leave_clause="Работнику предоставляется 30 календарных дней ежегодного основного оплачиваемого отпуска.",
        probation_clause="Условие об испытании сторонами не устанавливается.",
        function_clause="Работник оформляет приемы, переводы и увольнения, ведет кадровые документы и готовит сведения для обязательной отчетности.",
        structure_variant=4,
    ),
    ContractCase(
        filename="correct_05_legal_assistant.docx",
        category="correct",
        number="05-2026",
        city="Самара",
        date="15 сентября 2026 года",
        employer="ООО Правовой Ресурс",
        employer_representative="управляющего партнера Белова Артема Леонидовича",
        employee="Николаевой Дарьи Игоревны",
        position="помощника юриста",
        department="юридический отдел",
        workplace="г. Самара, ул. Молодогвардейская, д. 120",
        start_date="16 сентября 2026 года",
        salary_clause="Работнику устанавливается должностной оклад в размере 65 000 рублей в месяц. Премии могут выплачиваться по правилам положения о премировании.",
        payment_clause="Выплата заработной платы производится дважды в месяц: 17 числа текущего месяца и 2 числа следующего месяца на счет, выбранный Работником.",
        schedule_clause="Рабочая неделя является пятидневной с выходными днями в субботу и воскресенье. Рабочее время - с 9 часов 00 минут до 18 часов 00 минут, перерыв - с 13 часов 00 минут до 14 часов 00 минут.",
        hours_clause="Продолжительность рабочей недели составляет 40 часов.",
        leave_clause="Ежегодный основной оплачиваемый отпуск предоставляется на 28 календарных дней.",
        probation_clause="Испытательный срок составляет один месяц.",
        function_clause="Работник готовит проекты документов, систематизирует судебные материалы и выполняет правовой поиск по поручению руководителя.",
        structure_variant=5,
    ),
    ContractCase(
        filename="problematic_01_multiple_violations.docx",
        category="problematic",
        number="П-01-2026",
        city="Москва",
        date="18 сентября 2026 года",
        employer="ООО Быстрый Результат",
        employer_representative="директора Громова Олега Петровича",
        employee="Федорова Максима Олеговича",
        position="специалиста отдела продаж",
        department="отдел продаж",
        workplace="г. Москва, ул. Промышленная, д. 7",
        start_date="21 сентября 2026 года",
        salary_clause="Работнику устанавливается оклад 75 000 рублей в месяц.",
        payment_clause="Заработная плата выплачивается один раз в месяц 20 числа.",
        schedule_clause="Работник трудится пять дней в неделю с 9 часов 00 минут до 19 часов 00 минут с перерывом продолжительностью один час.",
        hours_clause="Продолжительность рабочего времени составляет 45 часов в неделю.",
        leave_clause="Ежегодный основной оплачиваемый отпуск предоставляется продолжительностью 14 календарных дней.",
        probation_clause="Работнику устанавливается испытательный срок шесть месяцев.",
        function_clause="Работник ведет переговоры с клиентами, готовит коммерческие предложения и оформляет сделки.",
    ),
    ContractCase(
        filename="problematic_02_vague_salary_and_long_week.docx",
        category="problematic",
        number="П-02-2026",
        city="Казань",
        date="19 сентября 2026 года",
        employer="ООО Регион Проект",
        employer_representative="генерального директора Власовой Ольги Игоревны",
        employee="Козлова Андрея Ильича",
        position="координатора проектов",
        department="проектный офис",
        workplace="г. Казань, ул. Баумана, д. 22",
        start_date="22 сентября 2026 года",
        salary_clause="Оплата труда определяется действующим штатным расписанием Работодателя. Конкретный размер оклада в настоящем договоре не устанавливается.",
        payment_clause="Заработная плата перечисляется в последний рабочий день каждого месяца.",
        schedule_clause="Рабочая неделя включает пять рабочих дней по десять часов с понедельника по пятницу. Перерыв - один час.",
        hours_clause="Общая продолжительность рабочего времени составляет 50 часов в неделю.",
        leave_clause="Работнику предоставляется ежегодный оплачиваемый отпуск на 20 календарных дней.",
        probation_clause="Срок испытания составляет три месяца.",
        function_clause="Работник координирует сроки проектов, собирает статусы и готовит отчеты для руководителя.",
        structure_variant=2,
    ),
    ContractCase(
        filename="problematic_03_probation_leave_and_bank.docx",
        category="problematic",
        number="П-03-2026",
        city="Екатеринбург",
        date="20 сентября 2026 года",
        employer="ООО Деловая Линия",
        employer_representative="директора Яковлева Романа Павловича",
        employee="Соколовой Ксении Артемовны",
        position="секретаря",
        department="администрация",
        workplace="г. Екатеринбург, ул. Восточная, д. 14",
        start_date="23 сентября 2026 года",
        salary_clause="Размер должностного оклада Работника составляет 60 000 рублей в месяц.",
        payment_clause="Заработная плата выплачивается 5 и 29 числа только на счет в банке, выбранном Работодателем. Изменение банка Работником не допускается.",
        schedule_clause="Рабочие дни - с понедельника по субботу. Время работы определяется графиком Работодателя, перерыв составляет один час.",
        hours_clause="Работнику устанавливается 42-часовая рабочая неделя.",
        leave_clause="Ежегодный основной оплачиваемый отпуск составляет 21 календарный день.",
        probation_clause="Испытание устанавливается на пять месяцев.",
        function_clause="Работник принимает корреспонденцию, организует встречи и ведет календарь руководителя.",
        structure_variant=3,
    ),
    ContractCase(
        filename="problematic_04_employer_discretion.docx",
        category="problematic",
        number="П-04-2026",
        city="Новосибирск",
        date="21 сентября 2026 года",
        employer="ООО Вектор Плюс",
        employer_representative="директора Зайцева Дениса Евгеньевича",
        employee="Лебедевой Ирины Алексеевны",
        position="маркетолога",
        department="отдел маркетинга",
        workplace="г. Новосибирск, ул. Советская, д. 31",
        start_date="24 сентября 2026 года",
        salary_clause="Работодатель определяет размер заработной платы по своему усмотрению и вправе изменить его без письменного соглашения с Работником.",
        payment_clause="Заработная плата выдается один раз в месяц 15 числа наличными в кассе Работодателя.",
        schedule_clause="Работник выполняет обязанности шесть дней в неделю по восемь часов. Выходной день - воскресенье.",
        hours_clause="Продолжительность рабочего времени составляет 48 часов в неделю.",
        leave_clause="Работнику предоставляется ежегодный основной оплачиваемый отпуск продолжительностью 14 календарных дней.",
        probation_clause="Испытательный срок составляет четыре месяца.",
        function_clause="Работник планирует рекламные кампании, готовит материалы и анализирует результаты продвижения.",
        structure_variant=4,
    ),
    ContractCase(
        filename="problematic_05_missing_terms.docx",
        category="problematic",
        number="П-05-2026",
        city="Самара",
        date="22 сентября 2026 года",
        employer="ООО Офис Сервис",
        employer_representative="директора Крылова Сергея Антоновича",
        employee="Поповой Марины Львовны",
        position="специалиста",
        department="без указания структурного подразделения",
        workplace=None,
        start_date="25 сентября 2026 года",
        salary_clause="Вознаграждение выплачивается в размере, который Работодатель определяет после завершения каждого месяца.",
        payment_clause="Расчет с Работником производится каждые 30 календарных дней.",
        schedule_clause="Работодатель самостоятельно определяет начало и окончание рабочего дня и может изменить график без уведомления Работника.",
        hours_clause="Рабочая неделя составляет 44 часа.",
        leave_clause="Ежегодный основной оплачиваемый отпуск предоставляется на 27 календарных дней.",
        probation_clause="Работнику устанавливается испытание продолжительностью три месяца.",
        function_clause="Работник выполняет любые поручения Работодателя. Конкретный вид поручаемой работы договором не определяется.",
        structure_variant=5,
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def keep_with_next(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:keepNext")) is None:
        p_pr.append(OxmlElement("w:keepNext"))


def add_clause(doc: Document, clause_id: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Inches(0.3)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(f"{clause_id} ")
    run.bold = True
    paragraph.add_run(text)


def add_section_heading(doc: Document, number: int, title: str, variant: int) -> None:
    if variant == 2:
        text = f"Раздел {number} {title}"
    elif variant == 3:
        text = f"{number} {title.upper()}"
    else:
        text = f"{number} {title}"
    paragraph = doc.add_paragraph(text, style="Heading 1")
    keep_with_next(paragraph)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    heading = doc.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(4)
    heading._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def add_title_block(doc: Document, case: ContractCase) -> None:
    title = doc.add_paragraph("Трудовой договор", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)

    subtitle = doc.add_paragraph("Тестовый пример без реальных персональных данных")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(9)

    number = doc.add_paragraph(f"Номер {case.number}")
    number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    number.paragraph_format.space_after = Pt(8)

    metadata = doc.add_table(rows=1, cols=2)
    metadata.autofit = False
    metadata.columns[0].width = Inches(3.35)
    metadata.columns[1].width = Inches(3.35)
    metadata.cell(0, 0).text = case.city
    metadata.cell(0, 1).text = case.date
    metadata.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    metadata.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for cell in metadata.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 70, 90, 70, 90)
    set_table_borders(metadata, color="FFFFFF", size="0")

    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(8)
    intro.paragraph_format.space_after = Pt(8)
    intro.paragraph_format.line_spacing = 1.08
    intro.add_run(case.employer).bold = True
    intro.add_run(", именуемое в дальнейшем Работодатель, в лице ")
    intro.add_run(case.employer_representative)
    intro.add_run(", действующего на основании Устава, с одной стороны, и ")
    intro.add_run(case.employee).bold = True
    intro.add_run(", именуемого в дальнейшем Работник, с другой стороны, заключили настоящий трудовой договор.")


def add_requisites(doc: Document, case: ContractCase) -> None:
    add_section_heading(doc, 8, "Реквизиты и подписи сторон", case.structure_variant)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.35)
    table.columns[1].width = Inches(3.35)
    employer = table.cell(0, 0)
    employee = table.cell(0, 1)
    employer.text = (
        f"Работодатель\n{case.employer}\nИНН 7700000000\n"
        "Адрес 000000 Российская Федерация\n\n______________ Представитель"
    )
    employee.text = (
        f"Работник\n{case.employee}\nПаспорт 00 00 000000\n"
        "Телефон +7 900 000 00 00\n\n______________ Работник"
    )
    for index, cell in enumerate(table.rows[0].cells):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 130, 150, 130, 150)
        set_cell_shading(cell, "F4F6F8" if index == 0 else "FFFFFF")
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.05
    set_table_borders(table)


def build_contract(case: ContractCase) -> Path:
    doc = Document()
    configure_styles(doc)
    add_title_block(doc, case)

    add_section_heading(doc, 1, "Предмет договора", case.structure_variant)
    add_clause(doc, "1.1", f"Работодатель принимает Работника на должность {case.position} в подразделение {case.department}.")
    add_clause(doc, "1.2", case.function_clause)
    if case.workplace:
        add_clause(doc, "1.3", f"Место работы: {case.workplace}.")
    else:
        add_clause(doc, "1.3", "Место выполнения работы в настоящем договоре не указывается и определяется Работодателем по мере необходимости.")
    add_clause(doc, "1.4", "Работа по настоящему договору является для Работника основной. Договор заключен на неопределенный срок.")
    add_clause(doc, "1.5", f"Дата начала работы: {case.start_date}.")

    add_section_heading(doc, 2, "Права и обязанности сторон", case.structure_variant)
    add_clause(doc, "2.1", "Работник обязан добросовестно исполнять трудовую функцию, соблюдать правила внутреннего трудового распорядка и требования охраны труда.")
    add_clause(doc, "2.2", "Работодатель обязан предоставить обусловленную договором работу, обеспечить безопасные условия труда и своевременно выплачивать заработную плату.")

    add_section_heading(doc, 3, "Оплата труда", case.structure_variant)
    add_clause(doc, "3.1", case.salary_clause)
    add_clause(doc, "3.2", case.payment_clause)
    add_clause(doc, "3.3", "При выплате заработной платы Работодатель предоставляет Работнику расчетный листок со сведениями о начислениях и удержаниях.")

    add_section_heading(doc, 4, "Рабочее время и время отдыха", case.structure_variant)
    add_clause(doc, "4.1", case.hours_clause)
    add_clause(doc, "4.2", case.schedule_clause)
    add_clause(doc, "4.3", case.leave_clause)

    add_section_heading(doc, 5, "Испытание", case.structure_variant)
    add_clause(doc, "5.1", case.probation_clause)

    add_section_heading(doc, 6, "Условия труда и страхование", case.structure_variant)
    add_clause(doc, "6.1", "Условия труда на рабочем месте являются допустимыми. Работа не имеет разъездного, подвижного или иного особого характера.")
    add_clause(doc, "6.2", "Работник подлежит обязательному социальному страхованию в соответствии с законодательством Российской Федерации.")

    add_section_heading(doc, 7, "Заключительные положения", case.structure_variant)
    add_clause(doc, "7.1", "Изменение условий настоящего договора оформляется письменным соглашением сторон, если законодательством не предусмотрено иное.")
    add_clause(doc, "7.2", "Договор составлен в двух экземплярах равной юридической силы. Один экземпляр передан Работнику, второй хранится у Работодателя.")

    add_requisites(doc, case)

    output_dir = OUTPUT_ROOT / case.category
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / case.filename
    doc.core_properties.title = "Трудовой договор"
    doc.core_properties.subject = "Тестовый документ для проверки анализатора"
    doc.core_properties.author = "Contract Checker MVP"
    doc.save(output_path)
    return output_path


def main() -> None:
    paths = [build_contract(case) for case in CASES]
    print(f"Создано документов: {len(paths)}")
    for path in paths:
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
