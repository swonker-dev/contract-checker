from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
CONTRACTS = ROOT / "data" / "contracts"


EXPECTED_PROBLEMS = {
    11: (r"45 часов", r"один раз в месяц"),
    12: (r"15 календарных дней", r"шесть месяцев", r"штраф 35 000"),
    13: (r"в пределах Новосибирской области", r"любые поручения"),
    14: (r"согласно штатному расписанию", r"смена кредитной организации.*не допускается"),
    15: (r"до 31 декабря 2027", r"задержать выплату"),
    16: (r"48 часов", r"не считаются сверхурочной"),
    17: (r"в возрасте 17 лет", r"испытание продолжительностью два месяца"),
    18: (r"пять месяцев", r"оплате не подлежит"),
    19: (r"без установленных дат", r"выходные.*не устанавливаются"),
    20: (r"одним платежом", r"21 календарный день", r"120 календарных дней"),
}


def all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def check_docx(path: Path) -> dict:
    errors: list[str] = []
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member:
            errors.append(f"damaged ZIP member: {bad_member}")
        names = set(archive.namelist())
        for required in ("word/document.xml", "word/styles.xml", "word/footer1.xml"):
            if required not in names:
                errors.append(f"missing {required}")
        footer_xml = archive.read("word/footer1.xml").decode("utf-8") if "word/footer1.xml" in names else ""
        if "PAGE" not in footer_xml or "NUMPAGES" not in footer_xml:
            errors.append("missing page fields")

    doc = Document(path)
    text = all_text(doc)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    if not doc.paragraphs or doc.paragraphs[0].text != "Трудовой договор":
        errors.append("missing title")
    if len(headings) != 10:
        errors.append(f"expected 10 headings, found {len(headings)}")
    if len(doc.tables) != 2:
        errors.append(f"expected 2 tables, found {len(doc.tables)}")
    if len(text) < 4_500:
        errors.append(f"document text too short: {len(text)}")
    if "[[" in text or "TODO" in text or "placeholder" in text.lower():
        errors.append("placeholder text found")
    if not any(heading.endswith("Реквизиты и подписи сторон") for heading in headings):
        errors.append("signature section missing")
    section = doc.sections[0]
    width_in = section.page_width / 914400
    height_in = section.page_height / 914400
    if abs(width_in - 8.5) > 0.02 or abs(height_in - 11.0) > 0.02:
        errors.append(f"unexpected page size: {width_in:.2f}x{height_in:.2f}")

    match = re.search(r"_(\d{2})_", path.name)
    case_number = int(match.group(1)) if match else -1
    if path.parent.name == "problematic":
        for pattern in EXPECTED_PROBLEMS.get(case_number, ()):
            if not re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL):
                errors.append(f"expected defect pattern not found: {pattern}")
    else:
        forbidden = (
            r"45 часов",
            r"48 часов",
            r"один раз в месяц",
            r"отпуск.*(?:15|21) календар",
            r"испытание.*(?:четыре|пять|шесть) месяц",
        )
        for pattern in forbidden:
            if re.search(pattern, text, flags=re.IGNORECASE | re.DOTALL):
                errors.append(f"unexpected intentionally bad pattern: {pattern}")

    return {
        "path": str(path.relative_to(ROOT)),
        "bytes": path.stat().st_size,
        "characters": len(text),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "headings": len(headings),
        "sha256_text": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "errors": errors,
    }


def main() -> None:
    files = sorted(
        path
        for folder in (CONTRACTS / "correct", CONTRACTS / "problematic")
        for path in folder.glob("*.docx")
        if re.search(r"_(?:1[1-9]|20)_", path.name)
    )
    results = [check_docx(path) for path in files]
    hashes = [item["sha256_text"] for item in results]
    summary = {
        "file_count": len(files),
        "unique_text_count": len(set(hashes)),
        "error_count": sum(len(item["errors"]) for item in results),
        "files": results,
    }
    if len(files) != 20:
        summary["error_count"] += 1
        summary["collection_error"] = f"expected 20 files, found {len(files)}"
    if len(set(hashes)) != len(hashes):
        summary["error_count"] += 1
        summary["uniqueness_error"] = "duplicate full texts found"
    output = ROOT / ".docx-qa" / "independent-structural-report.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in summary.items() if k != "files"}, ensure_ascii=False, indent=2))
    for item in results:
        status = "OK" if not item["errors"] else "ERROR"
        print(f"{status} {item['path']} chars={item['characters']} errors={item['errors']}")
    raise SystemExit(1 if summary["error_count"] else 0)


if __name__ == "__main__":
    main()
