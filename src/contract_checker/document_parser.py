from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .models import ParsedDocument, TextBlock


class UnsupportedDocumentError(ValueError):
    pass


def _normalize(text: str) -> str:
    return " ".join(text.replace("\u00a0", " ").split())


def parse_docx(path: Path) -> ParsedDocument:
    document = Document(path)
    blocks: list[TextBlock] = []
    block_number = 1

    for paragraph in document.paragraphs:
        text = _normalize(paragraph.text)
        if not text:
            continue
        blocks.append(TextBlock(block_id=f"p{block_number}", text=text))
        block_number += 1

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            cells = [_normalize(cell.text) for cell in row.cells]
            text = " | ".join(cell for cell in cells if cell)
            if text:
                blocks.append(
                    TextBlock(
                        block_id=f"table-{table_number}-row-{row_number}",
                        text=text,
                    )
                )

    return ParsedDocument(filename=path.name, blocks=blocks)


def parse_pdf(path: Path) -> ParsedDocument:
    reader = PdfReader(path)
    blocks: list[TextBlock] = []
    block_number = 1

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        for raw_paragraph in page_text.splitlines():
            text = _normalize(raw_paragraph)
            if not text:
                continue
            blocks.append(
                TextBlock(
                    block_id=f"p{block_number}",
                    text=text,
                    page_number=page_number,
                )
            )
            block_number += 1

    return ParsedDocument(filename=path.name, blocks=blocks)


def parse_document(path: str | Path) -> ParsedDocument:
    resolved = Path(path)
    suffix = resolved.suffix.lower()
    if suffix == ".docx":
        return parse_docx(resolved)
    if suffix == ".pdf":
        return parse_pdf(resolved)
    raise UnsupportedDocumentError("Поддерживаются только файлы DOCX и PDF")

