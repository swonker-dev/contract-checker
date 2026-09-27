import asyncio
from pathlib import Path

from docx import Document

from contract_checker.analyzer import ContractAnalyzer
from contract_checker.embeddings import HashingEmbedder
from contract_checker.knowledge_base import KnowledgeBase
from contract_checker.llm_client import LLMResponseFormatError, MockLLMClient
from contract_checker.models import CheckStatus, LLMChunkResult, RegulationChunk


class BrokenLLMClient:
    async def check_chunk(self, chunk, regulations):
        raise LLMResponseFormatError("Некорректный JSON")


class RecordingLLMClient:
    def __init__(self) -> None:
        self.calls = 0
        self.seen_articles: list[tuple[str, str | None]] = []

    async def check_chunk(self, chunk, regulations):
        self.calls += 1
        self.seen_articles.extend(
            (chunk.chunk_id, item.regulation.article) for item in regulations
        )
        return [
            LLMChunkResult(
                chunk_id=chunk.chunk_id,
                status=CheckStatus.WARNING,
                source=regulations[0].regulation.citation,
                comment="Семантическая проверка выполнена.",
            )
        ]


class SemanticRequirementsLLMClient:
    async def check_chunk(self, chunk, regulations):
        return []

    async def check_required_terms(self, chunks, regulations):
        return []


def test_end_to_end_analysis_in_mock_mode() -> None:
    path = Path("tests/.temporary_contract.docx")
    document = Document()
    document.add_paragraph("Трудовой договор")
    document.add_paragraph("1.1 Работодатель принимает Работника на должность аналитика.")
    document.add_paragraph("2.1 Место работы: город Москва.")
    document.add_paragraph("3.1 Трудовая функция определяется должностной инструкцией.")
    document.add_paragraph("4.1 Оплата труда составляет 80000 рублей в месяц.")
    document.add_paragraph("5.1 Режим рабочего времени составляет 40 часов в неделю.")
    document.save(path)

    try:
        knowledge_base = KnowledgeBase(
            chunks=[
                RegulationChunk(
                    chunk_id="tk-57",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="57",
                    text="Обязательными являются условия оплаты труда.",
                )
            ],
            embedder=HashingEmbedder(),
        )
        analyzer = ContractAnalyzer(knowledge_base, MockLLMClient(), top_k=1)

        report = asyncio.run(analyzer.analyze(path))

        assert report.supported_document
        assert report.chunks_total >= 5
        assert all(result.status == "not_applicable" for result in report.results)
        assert "не является юридическим заключением" in report.markdown
    finally:
        path.unlink(missing_ok=True)


def test_llm_format_failure_is_not_reported_as_legal_warning() -> None:
    path = Path("tests/.temporary_broken_llm_contract.docx")
    document = Document()
    document.add_paragraph("Трудовой договор")
    document.add_paragraph("1.1 Работник принимается на должность аналитика.")
    document.add_paragraph("2.1 Место работы: город Москва.")
    document.add_paragraph("3.1 Трудовая функция: анализ данных.")
    document.add_paragraph("4.1 Оклад составляет 80000 рублей в месяц.")
    document.add_paragraph("5.1 Рабочее время составляет 40 часов в неделю.")
    document.save(path)

    try:
        knowledge_base = KnowledgeBase(
            chunks=[
                RegulationChunk(
                    chunk_id="tk-57",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="57",
                    text="Обязательными являются условия оплаты труда.",
                )
            ],
            embedder=HashingEmbedder(),
        )
        analyzer = ContractAnalyzer(knowledge_base, BrokenLLMClient(), top_k=1)

        report = asyncio.run(analyzer.analyze(path))
        errors = [result for result in report.results if result.status == CheckStatus.ERROR]

        assert errors
        assert all(not result.source for result in errors)
        assert "Ошибка обработки" in report.markdown
        assert "JSONDecodeError" not in report.markdown
    finally:
        path.unlink(missing_ok=True)


def test_rule_finding_skips_only_the_article_already_resolved() -> None:
    path = Path("tests/.temporary_rule_and_llm_contract.docx")
    document = Document()
    document.add_paragraph("Трудовой договор")
    document.add_paragraph(
        "4.1 Продолжительность рабочего времени составляет 45 часов в неделю."
    )
    document.save(path)

    try:
        law_quote = (
            "Нормальная продолжительность рабочего времени не может превышать "
            "40 часов в неделю."
        )
        knowledge_base = KnowledgeBase(
            chunks=[
                RegulationChunk(
                    chunk_id="tk-91",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="91",
                    text=law_quote,
                ),
                RegulationChunk(
                    chunk_id="tk-115",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="115",
                    text=(
                        "Ежегодный основной оплачиваемый отпуск "
                        "предоставляется продолжительностью 28 календарных дней."
                    ),
                ),
            ],
            embedder=HashingEmbedder(),
        )
        llm = RecordingLLMClient()
        analyzer = ContractAnalyzer(knowledge_base, llm, top_k=2)

        report = asyncio.run(analyzer.analyze(path))

        assert llm.calls > 0
        chunk_articles = {
            article for chunk_id, article in llm.seen_articles if chunk_id == "4.1"
        }
        assert "91" not in chunk_articles
        assert "115" in chunk_articles
        assert any(result.status == CheckStatus.VIOLATION for result in report.results)
        semantic_result = next(
            result
            for result in report.results
            if result.comment == "Семантическая проверка выполнена."
        )
        assert semantic_result.status == CheckStatus.NOT_APPLICABLE
    finally:
        path.unlink(missing_ok=True)


def test_semantic_required_terms_can_override_regex_candidates() -> None:
    path = Path("tests/.temporary_semantic_requirements_contract.docx")
    document = Document()
    document.add_paragraph("Соглашение между организацией и специалистом")
    document.add_paragraph("Профессиональная роль специалиста — аналитик данных.")
    document.add_paragraph("Основная площадка исполнения обязанностей — город Москва.")
    document.add_paragraph("Фиксированная ежемесячная выплата — 80 000 рублей.")
    document.add_paragraph("Обязанности выполняются с понедельника по пятницу с 9 до 18.")
    document.save(path)

    try:
        knowledge_base = KnowledgeBase(
            chunks=[
                RegulationChunk(
                    chunk_id="tk-57",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="57",
                    text="Обязательными являются условия оплаты труда.",
                ),
                RegulationChunk(
                    chunk_id="tk-100",
                    source_id="tk",
                    source_title="ТК РФ",
                    article="100",
                    text="Режим рабочего времени должен быть определён.",
                ),
            ],
            embedder=HashingEmbedder(),
        )
        analyzer = ContractAnalyzer(
            knowledge_base,
            SemanticRequirementsLLMClient(),
            top_k=1,
        )

        report = asyncio.run(analyzer.analyze(path))

        assert not any(
            result.chunk_id == "document" and result.status == CheckStatus.WARNING
            for result in report.results
        )
    finally:
        path.unlink(missing_ok=True)
